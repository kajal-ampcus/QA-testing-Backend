import io
import json
import uuid

from typing import Literal

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Query,
    UploadFile,
)
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from apps.api.dependencies import get_db_session
from schemas.form_test_data import (
    FormTestDataCreate,
    FormTestDataUpdate,
    FormTestDataResponse,
)
from infra.db.models.form_test_data import FormTestData
from infra.db.models.project import Project
from infra.db.models.test_case import TestCase


router = APIRouter(
    prefix="/projects/{project_id}/form-test-data",
    tags=["form-test-data"],
)

MAX_FILE_SIZE = 5 * 1024 * 1024
ALLOWED_TEST_TYPES = {"positive", "negative", "validation", "edge"}


async def ensure_project(session: AsyncSession, project_id: uuid.UUID):
    project = await session.get(Project, project_id)
    if project is None:
        raise HTTPException(404, "Project not found")
    return project


async def ensure_case(
    session: AsyncSession,
    project_id: uuid.UUID,
    test_case_id: uuid.UUID | None,
    exclude_id: uuid.UUID | None = None,
) -> None:
    if test_case_id is None:
        return
    case = await session.scalar(
        select(TestCase).where(
            TestCase.id == test_case_id,
            TestCase.project_id == project_id,
        )
    )
    if case is None:
        raise HTTPException(422, "The selected test case does not belong to this project")
    duplicate = await session.scalar(
        select(FormTestData.id).where(
            FormTestData.project_id == project_id,
            FormTestData.test_case_id == test_case_id,
            FormTestData.active.is_(True),
            FormTestData.id != exclude_id if exclude_id else FormTestData.id.is_not(None),
        )
    )
    if duplicate:
        raise HTTPException(409, "A form fixture is already assigned to this test case")


async def get_record(
    session: AsyncSession,
    project_id: uuid.UUID,
    data_id: uuid.UUID,
):
    result = await session.scalar(
        select(FormTestData).where(
            FormTestData.id == data_id,
            FormTestData.project_id == project_id,
            FormTestData.active.is_(True),
        )
    )
    if result is None:
        raise HTTPException(404, "Active form test data not found")
    return result


@router.post("", response_model=FormTestDataResponse, status_code=201)
async def create_form_test_data(
    project_id: uuid.UUID,
    body: FormTestDataCreate,
    session: AsyncSession = Depends(get_db_session),
):
    await ensure_project(session, project_id)
    await ensure_case(session, project_id, body.test_case_id)

    row = FormTestData(
        project_id=project_id,
        **body.model_dump(),
    )
    session.add(row)
    await session.commit()
    await session.refresh(row)
    return row


@router.get("", response_model=list[FormTestDataResponse])
async def list_form_test_data(
    project_id: uuid.UUID,
    form_group: str | None = None,
    test_type: Literal[
        "positive", "negative", "validation", "edge"
    ] | None = None,
    tag: list[str] | None = Query(default=None),
    test_case_id: uuid.UUID | None = None,
    session: AsyncSession = Depends(get_db_session),
):
    await ensure_project(session, project_id)

    query = select(FormTestData).where(
        FormTestData.project_id == project_id,
        FormTestData.active.is_(True),
    )

    if form_group:
        query = query.where(FormTestData.form_group == form_group)
    if test_case_id:
        query = query.where(FormTestData.test_case_id == test_case_id)

    if test_type:
        query = query.where(FormTestData.test_type == test_type)

    if tag:
        # PostgreSQL JSONB contains all supplied tags.
        query = query.where(FormTestData.tags.contains(tag))

    result = await session.scalars(
        query.order_by(FormTestData.created_at.desc())
    )
    return list(result.all())


@router.get("/{data_id}", response_model=FormTestDataResponse)
async def get_form_test_data(
    project_id: uuid.UUID,
    data_id: uuid.UUID,
    session: AsyncSession = Depends(get_db_session),
):
    await ensure_project(session, project_id)
    return await get_record(session, project_id, data_id)


@router.patch("/{data_id}", response_model=FormTestDataResponse)
async def update_form_test_data(
    project_id: uuid.UUID,
    data_id: uuid.UUID,
    body: FormTestDataUpdate,
    session: AsyncSession = Depends(get_db_session),
):
    await ensure_project(session, project_id)
    row = await get_record(session, project_id, data_id)

    changes = body.model_dump(exclude_unset=True)
    if "test_case_id" in changes:
        await ensure_case(session, project_id, changes["test_case_id"], data_id)

    for field, value in changes.items():
        if value is None and field not in {"expected_result", "test_case_id"}:
            raise HTTPException(
                422, f"{field} cannot be null"
            )
        setattr(row, field, value)

    await session.commit()
    await session.refresh(row)
    return row


@router.delete("/{data_id}", status_code=204)
async def deactivate_form_test_data(
    project_id: uuid.UUID,
    data_id: uuid.UUID,
    session: AsyncSession = Depends(get_db_session),
):
    await ensure_project(session, project_id)
    row = await get_record(session, project_id, data_id)
    row.active = False
    await session.commit()


@router.post("/import", status_code=200)
async def import_form_test_data(
    project_id: uuid.UUID,
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_db_session),
):
    await ensure_project(session, project_id)

    if not file.filename or not file.filename.lower().endswith(".xlsx"):
        raise HTTPException(400, "Upload an .xlsx file")

    content = await file.read(MAX_FILE_SIZE + 1)

    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(413, "Excel file exceeds the 5 MB limit")

    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
        )
        sheet = workbook.active

        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            raise HTTPException(400, "Excel file is empty")

        headers = [
            str(value).strip().lower() if value is not None else ""
            for value in headers
        ]

        required = {
            "label",
            "form_group",
            "test_type",
            "test_data",
        }

        if not required.issubset(set(headers)):
            raise HTTPException(
                400,
                {
                    "message": "Missing required Excel columns",
                    "required_columns": sorted(required),
                    "received_columns": headers,
                },
            )

        column_indexes = {
            name: headers.index(name) for name in set(headers)
            if name
        }

        imported = []
        errors = []
        seen = set()
        assigned_cases: set[uuid.UUID] = set()

        for row_number, values in enumerate(rows, start=2):
            if not any(value is not None for value in values):
                continue

            try:
                raw = {
                    name: values[index] if index < len(values) else None
                    for name, index in column_indexes.items()
                }

                label = str(raw.get("label") or "").strip()
                form_group = str(raw.get("form_group") or "").strip()
                test_type = str(raw.get("test_type") or "").strip().lower()

                if not label or not form_group:
                    raise ValueError("label and form_group are required")

                if test_type not in ALLOWED_TEST_TYPES:
                    raise ValueError(
                        f"test_type must be one of "
                        f"{sorted(ALLOWED_TEST_TYPES)}"
                    )

                raw_test_data = raw.get("test_data")

                if isinstance(raw_test_data, dict):
                    test_data = raw_test_data
                elif isinstance(raw_test_data, str):
                    test_data = json.loads(raw_test_data)
                else:
                    raise ValueError(
                        "test_data must be a JSON object in an Excel cell"
                    )

                if not isinstance(test_data, dict):
                    raise ValueError("test_data must be a JSON object")

                raw_tags = raw.get("tags")
                if raw_tags is None:
                    tags = []
                elif isinstance(raw_tags, str):
                    tags = [
                        item.strip()
                        for item in raw_tags.split(",")
                        if item.strip()
                    ]
                elif isinstance(raw_tags, list):
                    tags = raw_tags
                else:
                    raise ValueError("tags must be comma-separated text")

                expected = raw.get("expected_result")
                expected_result = (
                    str(expected).strip() if expected is not None else None
                )

                raw_case_id = raw.get("test_case_id")
                test_case_code = str(raw.get("test_case_code") or "").strip()
                if raw_case_id and test_case_code:
                    raise ValueError("Provide either test_case_id or test_case_code, not both")
                if test_case_code:
                    matched_case = await session.scalar(
                        select(TestCase).where(
                            TestCase.project_id == project_id,
                            TestCase.tc_code == test_case_code,
                        )
                    )
                    if matched_case is None:
                        raise ValueError(
                            f"No test case with code {test_case_code!r} exists in this project"
                        )
                    test_case_id = matched_case.id
                elif raw_case_id:
                    test_case_id = uuid.UUID(str(raw_case_id))
                else:
                    test_case_id = None

                payload = FormTestDataCreate(
                    label=label,
                    form_group=form_group,
                    test_case_id=test_case_id,
                    test_type=test_type,
                    test_data=test_data,
                    expected_result=expected_result,
                    tags=tags,
                )

                if payload.test_case_id:
                    await ensure_case(session, project_id, payload.test_case_id)
                    if payload.test_case_id in assigned_cases:
                        raise ValueError("Only one form fixture can be assigned to a test case")
                    assigned_cases.add(payload.test_case_id)

                duplicate_key = (
                    payload.form_group,
                    payload.label,
                )

                if duplicate_key in seen:
                    raise ValueError("Duplicate label and form_group in upload")

                seen.add(duplicate_key)

                imported.append(
                    FormTestData(
                        project_id=project_id,
                        **payload.model_dump(),
                    )
                )

            except HTTPException as exc:
                errors.append({"row": row_number, "error": str(exc.detail)})
            except HTTPException as exc:
                errors.append({"row": row_number, "error": str(exc.detail)})
            except (ValueError, TypeError, json.JSONDecodeError, ValidationError) as exc:
                errors.append({
                    "row": row_number,
                    "error": str(exc),
                })

        if imported:
            session.add_all(imported)
            await session.commit()

        return {
            "imported_count": len(imported),
            "error_count": len(errors),
            "errors": errors,
        }

    except HTTPException:
        raise
    except Exception as exc:
        await session.rollback()
        raise HTTPException(
            400, "Could not read the Excel file"
        ) from exc
    finally:
        if "workbook" in locals():
            workbook.close()
        await file.close()
