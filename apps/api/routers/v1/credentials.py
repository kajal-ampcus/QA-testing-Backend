"""Project-scoped accounts. Passwords remain encrypted and are never returned."""
import io
import uuid
from datetime import datetime

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import AnyHttpUrl, BaseModel, Field, ValidationError, field_validator, model_validator
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession
from typing import Literal

from apps.api.dependencies import get_db_session
from infra.db.models.discovery_credential import DiscoveryCredential
from infra.db.models.application_map import ApplicationMap
from infra.db.models.project import Project
from infra.db.models.test_case import TestCase
from infra.secrets.vault_client import decrypt_json_secret, encrypt_login_secret
from core.agents.application_discovery.login_choices import login_group_key

router = APIRouter(prefix='/projects/{project_id}/credentials', tags=['credentials'])
MAX_IMPORT_BYTES = 5 * 1024 * 1024

class CredentialCreateRequest(BaseModel):
    label: str = Field(default='Test account',min_length=1,max_length=120)
    role: str = Field(default='User',min_length=1,max_length=120)
    username: str = Field(min_length=1,max_length=500)
    password: str = Field(min_length=1,max_length=500)
    login_url: AnyHttpUrl | None = None
    username_selector: str | None = Field(default=None,max_length=200)
    password_selector: str | None = Field(default=None,max_length=200)
    submit_selector: str | None = Field(default=None,max_length=200)
    test_type : Literal["positive","negative","edge"]= 'positive'
    credential_type: Literal["login"] = "login"
    tags: list[str] = Field(default_factory=list)
    set_as_project_default: bool = True

    @field_validator('label','role','username')
    @classmethod
    def nonblank(cls,value):
        if not value.strip():
            raise ValueError('Must not be blank')
        return value.strip()

class CredentialUpdateRequest(CredentialCreateRequest):
    username: str | None = Field(default=None,min_length=1,max_length=500)
    password: str | None = Field(default=None,min_length=1,max_length=500)
    set_as_project_default: bool = False
    test_type: Literal[ "positive", "negative", "edge"] | None = None
    credential_type: Literal["login"] | None = None
    tags: list[str] | None = None

    @field_validator('username')
    @classmethod
    def optional_username(cls,value):
        return value

    @model_validator(mode='after')
    def paired_credentials(self):
        if (self.username is None) != (self.password is None):
            raise ValueError('Supply both username and password to replace login credentials')
        return self

class CredentialCreateResponse(BaseModel):
    credential_ref: str
    label: str
    role: str
    is_default: bool
    created_at: datetime
    login_group: str | None = None
    login_route: str | None = None
    test_type: str
    credential_type: str
    tags: list[str]


def response(row,project):
    return CredentialCreateResponse(credential_ref=row.credential_ref,label=row.label,role=row.role,is_default=project.credential_ref==row.credential_ref,created_at=row.created_at,login_group=row.login_group,login_route=row.login_route, test_type=row.test_type, credential_type=row.credential_type, tags=row.tags or [],)

async def project_for(session,project_id):
    project=await session.get(Project,project_id)
    if project is None:
        raise HTTPException(404,'Project not found')
    return project

async def account_for(session,project_id,ref):
    row=await session.scalar(select(DiscoveryCredential).where(DiscoveryCredential.project_id==project_id,DiscoveryCredential.credential_ref==ref,DiscoveryCredential.active.is_(True)).with_for_update())
    if row is None:
        raise HTTPException(404,'Active account not found in this project')
    return row

def secret_for(body):
    secret=body.model_dump(mode='json',include={'username','password','login_url','username_selector','password_selector','submit_selector'},exclude_none=True)
    try:
        return encrypt_login_secret(secret)
    except RuntimeError as exc:
        raise HTTPException(503,'Credential encryption is not configured. Configure CREDENTIAL_ENCRYPTION_KEY on the backend.') from exc

def credential_identity(row, label: str, role: str, login_group: str, test_type: str, username: str) -> bool:
    """Match the same saved login account without ever exposing secret values."""
    if (
        row.label.strip().casefold() != label.strip().casefold()
        or row.role.strip().casefold() != role.strip().casefold()
        or row.login_group != login_group
        or row.test_type != test_type
        or row.credential_type != "login"
    ):
        return False
    try:
        saved = decrypt_json_secret(row.encrypted_secret)
    except RuntimeError:
        return False
    return str(saved.get("username") or "").strip().casefold() == username.strip().casefold()

async def upsert_credential(session: AsyncSession, project: Project, body: CredentialCreateRequest) -> DiscoveryCredential:
    route = str(body.login_url or project.application_url or "")
    group = login_group_key(route, body.role)
    rows = list(await session.scalars(
        select(DiscoveryCredential)
        .where(DiscoveryCredential.project_id == project.id, DiscoveryCredential.active.is_(True))
        .order_by(DiscoveryCredential.created_at)
        .with_for_update()
    ))
    matches = [row for row in rows if credential_identity(
        row, body.label, body.role, group, body.test_type, body.username
    )]
    row = next((item for item in matches if item.credential_ref == project.credential_ref), None) or (matches[0] if matches else None)
    if row is None:
        row = DiscoveryCredential(
            project_id=project.id,
            credential_ref=f"cred:{uuid.uuid4()}",
            label=body.label,
            role=body.role,
            encrypted_secret=secret_for(body),
            test_type=body.test_type,
            credential_type=body.credential_type,
            tags=body.tags,
        )
        row.login_route = route
        row.login_group = group
        session.add(row)
    else:
        row.label = body.label
        row.role = body.role
        row.encrypted_secret = secret_for(body)
        row.test_type = body.test_type
        row.credential_type = body.credential_type
        row.tags = body.tags
        row.login_route = route
        row.login_group = group

    # Repair legacy duplicate active rows for this exact login identity while
    # preserving all references and keeping the saved secret in encrypted form.
    duplicates = [item for item in matches if item is not row]
    duplicate_refs = [item.credential_ref for item in duplicates]
    if duplicate_refs:
        await session.execute(
            update(TestCase)
            .where(TestCase.project_id == project.id, TestCase.credential_ref.in_(duplicate_refs))
            .values(credential_ref=row.credential_ref)
        )
        maps = await session.scalars(select(ApplicationMap).where(ApplicationMap.project_id == project.id))
        for app_map in maps:
            if app_map.credential_ref in duplicate_refs:
                app_map.credential_ref = row.credential_ref
            checkpoint = app_map.discovery_checkpoint
            configuration = checkpoint.get("configuration") if checkpoint else None
            if configuration and configuration.get("credential_ref") in duplicate_refs:
                configuration["credential_ref"] = row.credential_ref
                app_map.discovery_checkpoint = checkpoint
        for duplicate in duplicates:
            duplicate.active = False
        if project.credential_ref in duplicate_refs:
            project.credential_ref = row.credential_ref

    if body.set_as_project_default:
        project.credential_ref = row.credential_ref
    return row

@router.get('',response_model=list[CredentialCreateResponse])
async def list_credentials(project_id:uuid.UUID,session:AsyncSession=Depends(get_db_session)):
    project=await project_for(session,project_id)
    rows=await session.scalars(select(DiscoveryCredential).where(DiscoveryCredential.project_id==project_id,DiscoveryCredential.active.is_(True)).order_by(DiscoveryCredential.created_at))
    return [response(row,project) for row in rows]

@router.get('/template')
async def credential_import_template(project_id: uuid.UUID, session: AsyncSession = Depends(get_db_session)):
    """Download the Excel template accepted by the credential import endpoint."""
    await project_for(session, project_id)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Login credentials'
    sheet.append(['label', 'role', 'username', 'password', 'login_url', 'test_type', 'tags'])
    sheet.append(['Admin account', 'Admin', 'admin@example.com', 'replace-me', '', 'positive', 'admin, smoke'])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = 'A1:G2'
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return StreamingResponse(
        output,
        media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename="login-credentials-template.xlsx"'},
    )

@router.post('',response_model=CredentialCreateResponse,status_code=201)
async def create_credential(project_id:uuid.UUID,body:CredentialCreateRequest,session:AsyncSession=Depends(get_db_session)):
    project=await project_for(session,project_id)
    row=await upsert_credential(session, project, body)
    await session.commit()
    await session.refresh(row)
    return response(row,project)

@router.post('/{credential_ref}/revisions',response_model=CredentialCreateResponse,status_code=201)
async def revise_credential(project_id:uuid.UUID,credential_ref:str,body:CredentialUpdateRequest,session:AsyncSession=Depends(get_db_session)):
    project=await project_for(session,project_id)
    old=await account_for(session,project_id,credential_ref)
    row=DiscoveryCredential(
        project_id=project_id,
        credential_ref=f'cred:{uuid.uuid4()}',
        label=body.label,
        role=body.role,
        encrypted_secret=secret_for(body) if body.username is not None else old.encrypted_secret,
        test_type=body.test_type or old.test_type,
        credential_type=body.credential_type or old.credential_type,
        tags=body.tags if body.tags is not None else (old.tags or []),
    )
    row.login_route = str(body.login_url or old.login_route or project.application_url or "")
    row.login_group = login_group_key(row.login_route, body.role)
    old.active=False
    session.add(row)
    await session.execute(
        update(TestCase)
        .where(TestCase.project_id == project_id, TestCase.credential_ref == old.credential_ref)
        .values(credential_ref=row.credential_ref)
    )
    if body.set_as_project_default:
        project.credential_ref=row.credential_ref
    elif project.credential_ref == old.credential_ref:
        project.credential_ref=None
    await session.commit()
    await session.refresh(row)
    return response(row,project)

@router.delete('/{credential_ref}',status_code=204)
async def delete_credential(project_id:uuid.UUID,credential_ref:str,session:AsyncSession=Depends(get_db_session)) -> None:
    """Deactivate an account. The encrypted row is kept for audit but can no longer be resolved."""
    project=await project_for(session,project_id)
    row=await account_for(session,project_id,credential_ref)
    assigned = await session.scalar(
        select(TestCase.id).where(
            TestCase.project_id == project_id,
            TestCase.credential_ref == credential_ref,
        ).limit(1)
    )
    if assigned is not None:
        raise HTTPException(
            409,
            "This credential is assigned to test cases. Reassign those cases before deactivating it.",
        )
    row.active=False
    if project.credential_ref==row.credential_ref:
        project.credential_ref=None
    await session.commit()

@router.post('/{credential_ref}/default',response_model=CredentialCreateResponse)
async def set_default_credential(project_id:uuid.UUID,credential_ref:str,session:AsyncSession=Depends(get_db_session)):
    project=await project_for(session,project_id)
    row=await account_for(session,project_id,credential_ref)
    project.credential_ref=row.credential_ref
    await session.commit()
    return response(row,project)


@router.post('/import', status_code=200)
async def import_credentials(
    project_id: uuid.UUID,
    file: UploadFile = File(...),
    session: AsyncSession = Depends(get_db_session),
):
    """Import .xlsx login rows; secrets are encrypted and never returned."""
    project = await project_for(session, project_id)
    if not file.filename or not file.filename.lower().endswith('.xlsx'):
        raise HTTPException(400, 'Upload an .xlsx file')
    content = await file.read(MAX_IMPORT_BYTES + 1)
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(413, 'Excel file exceeds the 5 MB limit')
    imported_count = 0
    errors: list[dict[str, str | int]] = []
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        first = next(rows, None)
        if not first:
            raise HTTPException(400, 'Excel file is empty')
        headers = [str(value).strip().lower() if value is not None else '' for value in first]
        required = {'label', 'role', 'username', 'password'}
        if not required.issubset(headers):
            raise HTTPException(400, {
                'message': 'Missing required Excel columns',
                'required_columns': sorted(required),
                'optional_columns': ['login_url', 'test_type', 'tags'],
            })
        indexes = {name: headers.index(name) for name in set(headers) if name}
        seen: set[tuple[str, str, str, str, str]] = set()
        for row_number, values in enumerate(rows, start=2):
            if not any(value is not None for value in values):
                continue
            try:
                raw = {name: values[i] if i < len(values) else None for name, i in indexes.items()}
                label = str(raw.get('label') or '').strip()
                role = str(raw.get('role') or '').strip()
                username = str(raw.get('username') or '').strip()
                password = str(raw.get('password') or '')
                if not label or not role or not username or not password:
                    raise ValueError('label, role, username and password are required')
                test_type = str(raw.get('test_type') or 'positive').strip().lower()
                if test_type not in {'positive', 'negative', 'edge'}:
                    raise ValueError('test_type must be positive, negative or edge')
                tags_value = raw.get('tags')
                tags = [part.strip() for part in str(tags_value or '').split(',') if part.strip()]
                body = CredentialCreateRequest(
                    label=label,
                    role=role,
                    username=username,
                    password=password,
                    login_url=str(raw['login_url']).strip() if raw.get('login_url') else None,
                    test_type=test_type,
                    tags=tags,
                    set_as_project_default=False,
                )
                route = str(body.login_url or project.application_url or '')
                key = (label.casefold(), role.casefold(), login_group_key(route, role), test_type, username.casefold())
                if key in seen:
                    raise ValueError('Duplicate login account identity in upload')
                seen.add(key)
                saved = await upsert_credential(session, project, body)
                if not project.credential_ref:
                    project.credential_ref = saved.credential_ref
                imported_count += 1
            except ValidationError as exc:
                messages = "; ".join(str(item["msg"]) for item in exc.errors())
                errors.append({'row': row_number, 'error': messages})
            except (ValueError, TypeError) as exc:
                errors.append({'row': row_number, 'error': str(exc)})
        if imported_count:
            await session.commit()
        return {'imported_count': imported_count, 'error_count': len(errors), 'errors': errors}
    except HTTPException:
        raise
    except Exception as exc:
        await session.rollback()
        raise HTTPException(400, 'Could not read the Excel file') from exc
    finally:
        if 'workbook' in locals():
            workbook.close()
        await file.close()
